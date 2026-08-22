#!/usr/bin/env python3
"""Build an A4 academic thesis (Word) from Thesis_Complete.md."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
MD = ROOT / "Thesis_Complete.md"
OUT = ROOT / "Thesis_Complete.docx"
HEADER_IMG = ROOT / "assets" / "cover" / "protocol_header.png"
FONT = "Times New Roman"


def _set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def _ppr(p):
    pPr = p._p.get_or_add_pPr()
    return pPr


def _set_spacing(p, *, before=0, after=0, double=True, first_line=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                 keep_with_next=False, page_break_before=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if double:
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.line_spacing = 2.0
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.line_spacing = 1.0
    pf.first_line_indent = Cm(first_line) if first_line else Cm(0)
    pf.widow_control = True
    pf.keep_with_next = keep_with_next
    pf.page_break_before = page_break_before


def _enable_update_fields(doc):
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


def _add_page_number(paragraph, fmt="decimal"):
    """Insert PAGE field."""
    run = paragraph.add_run()
    _set_run_font(run, 12)
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGE \\* {fmt} "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instr)
    run._r.append(fldChar2)


def _setup_footer(section, fmt="decimal", restart=None):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.clear()
    _set_spacing(p, double=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, before=0)
    _add_page_number(p, fmt)
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:fmt"), "lowerRoman" if fmt == "roman" else "decimal")
    if restart is not None:
        pgNumType.set(qn("w:start"), str(restart))


def _hide_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.clear()


def _style_heading(style, size, bold=True, align="center", space_before=0, space_after=12):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{a}"), FONT)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.line_spacing = 2.0
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Cm(0)
    if align == "center":
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{a}"), FONT)
    np = normal.paragraph_format
    np.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    np.line_spacing = 2.0
    np.space_before = Pt(0)
    np.space_after = Pt(0)
    np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _style_heading(doc.styles["Heading 1"], 14, True, "center", 0, 6)
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    _style_heading(doc.styles["Heading 2"], 12, True, "left", 18, 6)
    _style_heading(doc.styles["Heading 3"], 12, True, "left", 12, 6)


def centered(doc, text, size=12, bold=False, italic=False, before=0, after=0, double=True):
    p = doc.add_paragraph()
    _set_spacing(p, before=before, after=after, double=double, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def body(doc, text, first=True):
    p = doc.add_paragraph()
    _set_spacing(p, first_line=1.27 if first else 0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    render_inline(p, text)
    return p


def render_inline(p, text, size=12):
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        _set_run_font(run, size=size, bold=(idx % 2 == 1))


def add_heading1(doc, text, page_break=True):
    p = doc.add_paragraph(text, style="Heading 1")
    _set_spacing(
        p,
        before=0,
        after=6,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
        page_break_before=page_break,
    )
    for run in p.runs:
        _set_run_font(run, 14, bold=True)
    return p


def add_heading2(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    _set_spacing(p, before=18, after=6, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True, first_line=0)
    for run in p.runs:
        _set_run_font(run, 12, bold=True)
    return p


def shade_cell(cell, hex_color="E8E8E8"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            _set_spacing(p, double=False, after=2, before=2, align=WD_ALIGN_PARAGRAPH.LEFT)
            text = row[c_idx] if c_idx < len(row) else ""
            run = p.add_run(text)
            _set_run_font(run, size=9, bold=(r_idx == 0))
            set_cell_border(cell)
            if r_idx == 0:
                shade_cell(cell)
    # thin spacer so caption/table stay together visually
    sp = doc.add_paragraph()
    _set_spacing(sp, double=False, after=8, before=0)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if raw and all(set(c) <= set("-: ") and c for c in raw):
            i += 1
            continue
        rows.append(raw)
        i += 1
    return rows, i


def add_protocol_header(doc):
    """Bilingual Faculty of Dentistry header cropped from protocol page 1."""
    p = doc.add_paragraph()
    _set_spacing(p, double=False, before=0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    # Content width: 21 − 3.0 − 2.5 cm
    run.add_picture(str(HEADER_IMG), width=Cm(15.5))
    return p


def _set_run_arabic(run, size=14, bold=True):
    arabic = "Noto Naskh Arabic"
    run.font.name = arabic
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{a}"), arabic)
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)


def centered_ar(doc, text, size=14, bold=True, before=0, after=6):
    p = doc.add_paragraph()
    _set_spacing(p, before=before, after=after, double=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    run = p.add_run(text)
    _set_run_arabic(run, size=size, bold=bold)
    return p


def add_cover(doc):
    add_protocol_header(doc)
    centered(doc, "Faculty of Dentistry", 12, True, after=0, double=False)
    centered(doc, "Department of Conservative Dentistry", 12, italic=True, after=6, double=False)
    centered(doc, "Student Code No. 202203112", 12, True, after=6, double=False)
    centered(doc, "A Thesis submitted in partial fulfillment of the", 12, double=False)
    centered(doc, "requirements for the degree of Master of Science", 12, double=False)
    centered(doc, "in Conservative Dentistry", 12, True, after=6, double=False)
    centered(doc, "Academic Year 2024–2025", 12, True, after=6, double=False)
    centered(doc, "Name of Candidate: Mohamed Talaat Mohamed AbdelMoaty ElAbd", 12, True, after=10, double=False)

    centered(doc, "English Title", 12, True, after=0, double=False)
    centered(
        doc,
        "COMPARATIVE STUDY OF WEAR RESISTANCE AND SURFACE ROUGHNESS OF INJECTABLE VERSUS CONVENTIONAL COMPOSITE RESIN — IN VITRO STUDY",
        13,
        True,
        before=0,
        after=8,
        double=False,
    )
    centered(doc, "Arabic Title", 12, True, after=2, double=False)
    centered_ar(
        doc,
        "دراسة مقارنة للتآكل و خشونة سطح الراتينج المركب القابل للحقن و التقليدي - دراسة في المختبر",
        14,
        True,
        before=0,
        after=10,
    )

    centered(doc, "Supervision Committee", 12, True, before=4, after=4, double=False)
    centered(doc, "1. Prof. Wegdan M. Abdel-Fattah", 12, True, after=0, double=False)
    centered(doc, "2. Asst. Prof. Emad M. El-Sayed (Main supervisor)", 12, True, after=0, double=False)


def add_supervisors_page(doc):
    add_heading1(doc, "SUPERVISION COMMITTEE", page_break=False)
    body(doc, "This thesis was carried out under the supervision of:", first=False)
    centered(doc, "1. Prof. Wegdan M. Abdel-Fattah", 12, True, before=12)
    centered(doc, "Professor of Conservative Dentistry", 12, italic=True, after=12)
    centered(doc, "2. Asst. Prof. Emad M. El-Sayed", 12, True, before=12)
    centered(doc, "Assistant Professor of Conservative Dentistry", 12, italic=True)
    centered(doc, "(Main supervisor)", 12, italic=True)
    body(doc, "Faculty of Dentistry, Pharos University in Alexandria.", first=False)


def add_declaration(doc):
    add_heading1(doc, "DECLARATION", page_break=True)
    body(
        doc,
        "I declare that this thesis is my own work. It has not been submitted, in whole or in part, "
        "for any other degree or qualification. The experimental data are those of the laboratory "
        "study described herein. Sources used in the text are cited in the list of references.",
    )
    centered(doc, "Mohamed Talaat Mohamed AbdelMoaty ElAbd", 12, True, before=24)
    centered(doc, "Signature: …………………………     Date: …………………………", 12, before=12)


def add_abstract(doc):
    add_heading1(doc, "ABSTRACT", page_break=True)
    blocks = [
        (
            "Background. ",
            "Highly filled injectable composites are sold for load-bearing use. Their loss of mass and change in surface roughness after toothbrush abrasion still need a direct comparison with a conventional nanohybrid resin under the same laboratory conditions.",
        ),
        (
            "Aim. ",
            "To compare percentage weight loss and arithmetic mean roughness (Ra) of Beautifil Flow Plus X F00, G-ænial Universal Injectable and Beautifil II LS after 10 000 toothbrushing cycles.",
        ),
        (
            "Methods. ",
            "Thirty-six discs (10 mm diameter × 1 mm thick; n = 12) were brushed at 2 N in a Colgate Total slurry prepared after ISO guidance. Mass was recorded before and after abrasion. Ra was measured by contact profilometry. Percentage weight loss was compared by one-way ANOVA. Ra after brushing was compared by the Kruskal–Wallis test (α = 0.05).",
        ),
        (
            "Results. ",
            "Mean percentage weight loss was 1.49 ± 3.14 % for Beautifil Flow Plus X F00, 0.65 ± 0.37 % for G-ænial Universal Injectable and 0.41 ± 0.25 % for Beautifil II LS (p = 0.329). Mean Ra after brushing was 0.104 ± 0.017 µm, 0.100 ± 0.023 µm and 0.194 ± 0.050 µm respectively (p < 0.001). Each injectable resin was smoother than Beautifil II LS. The two injectables did not differ from each other. One Beautifil Flow Plus X F00 disc lost about 11.4 % of its mass and accounts for that group’s large standard deviation.",
        ),
        (
            "Conclusions. ",
            "Under this toothbrushing protocol, gravimetric wear did not differ significantly among the three resins. Both injectable materials remained significantly smoother than the conventional nanohybrid. The injectable means stayed below the 0.2 µm plaque threshold. The nanohybrid mean sat on that line.",
        ),
        (
            "Keywords. ",
            "Injectable composite; nanohybrid composite; toothbrush abrasion; wear; surface roughness; giomer.",
        ),
    ]
    for lead, rest in blocks:
        p = doc.add_paragraph()
        _set_spacing(p, first_line=1.27)
        r = p.add_run(lead)
        _set_run_font(r, 12, bold=True)
        r2 = p.add_run(rest)
        _set_run_font(r2, 12)


def add_contents(doc, entries):
    add_heading1(doc, "CONTENTS", page_break=True)
    for kind, text, page in entries:
        p = doc.add_paragraph()
        indent = 0 if kind == "ch" else 1.0
        _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0)
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(f"{text}\t{page}")
        _set_run_font(run, 12, bold=(kind == "ch"))


def add_list_of_tables(doc):
    add_heading1(doc, "LIST OF TABLES", page_break=True)
    items = [
        "Table 3.1  Brand, type, matrix, filler and load of the resin composites used in the study",
        "Table 4.1  Mean weight before and after simulated toothbrushing and percentage weight loss",
        "Table 4.2  Mean surface roughness (Ra) before and after simulated toothbrushing and ΔRa",
        "Table 4.3  Absolute weight loss after 10 000 brushing cycles",
    ]
    for item in items:
        p = doc.add_paragraph()
        _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        run = p.add_run(item)
        _set_run_font(run, 12)


def add_abbreviations(doc):
    add_heading1(doc, "LIST OF ABBREVIATIONS", page_break=True)
    rows = [
        ("ANOVA", "Analysis of variance"),
        ("ISO", "International Organization for Standardization"),
        ("Ra", "Arithmetic mean roughness"),
        ("SD", "Standard deviation"),
        ("S-PRG", "Surface pre-reacted glass-ionomer"),
        ("ΔRa", "Change in Ra (after − before)"),
    ]
    for abbr, meaning in rows:
        p = doc.add_paragraph()
        _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(f"{abbr}")
        _set_run_font(r, 12, bold=True)
        r2 = p.add_run(f"    {meaning}")
        _set_run_font(r2, 12)


def convert_body(doc, text):
    lines = text.splitlines()
    toc = []
    i = 0
    n = len(lines)
    pending_chapter = None
    first_chapter = True
    in_refs = False
    while i < n:
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("**Table ") and line.endswith("**") is False:
            # caption line **Table x.x** rest
            p = doc.add_paragraph()
            _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True, before=12, after=6)
            render_inline(p, line)
            i += 1
            continue
        if line.startswith("# CHAPTER"):
            pending_chapter = line[2:].strip()
            i += 1
            continue
        if line.startswith("# ") and pending_chapter:
            title = line[2:].strip()
            add_heading1(doc, pending_chapter, page_break=not first_chapter)
            first_chapter = False
            add_heading1(doc, title, page_break=False)
            if "REFERENCES" in title:
                in_refs = True
            pending_chapter = None
            i += 1
            continue
        if line.startswith("# "):
            add_heading1(doc, line[2:].strip(), page_break=True)
            toc.append(("ch", line[2:].strip()))
            i += 1
            continue
        if line.startswith("## "):
            h = line[3:].strip()
            add_heading2(doc, h)
            toc.append(("sec", h))
            i += 1
            continue
        # numbered list / references
        if line[:3].split(".")[0].isdigit() and (line[1] == "." or line[2] == "."):
            p = doc.add_paragraph()
            if in_refs:
                _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, double=False, after=6)
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.75)
            else:
                _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            render_inline(p, line)
            i += 1
            continue
        # table note
        if line.startswith("†") or line.startswith("‡") or line.startswith("Values are") or line.startswith("Manufacturers:"):
            p = doc.add_paragraph()
            _set_spacing(p, first_line=0, align=WD_ALIGN_PARAGRAPH.LEFT, double=False, after=8)
            render_inline(p, line, size=10)
            i += 1
            continue
        body(doc, line, first=True)
        i += 1
    return toc


def set_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)


def main():
    doc = Document()
    configure_styles(doc)
    _enable_update_fields(doc)
    set_margins(doc.sections[0])
    doc.sections[0].top_margin = Cm(1.5)
    doc.sections[0].different_first_page_header_footer = True
    _hide_footer(doc.sections[0])

    add_cover(doc)

    # Front matter section — roman numerals
    sect = doc.add_section()
    set_margins(sect)
    _setup_footer(sect, fmt="roman", restart=1)

    md = MD.read_text(encoding="utf-8")
    chapter_pages = {
        "CHAPTER 1": 1,
        "1.1 Aim of the Study": 2,
        "1.2 Objectives": 2,
        "1.3 Null Hypotheses": 2,
        "CHAPTER 2": 3,
        "2.1 Wear of Restorative Resins": 3,
        "2.2 Surface Roughness": 4,
        "2.3 Injectable Composite Resins": 5,
        "2.4 Conventional Nanohybrid Resins and Giomers": 6,
        "2.5 Statement of the Problem": 6,
        "CHAPTER 3": 7,
        "3.1 Study Design and Sample": 7,
        "3.2 Materials": 7,
        "3.3 Specimen Preparation": 8,
        "3.4 Baseline Measurements": 8,
        "3.5 Toothbrushing Protocol": 8,
        "3.6 Post-test Evaluation": 8,
        "3.7 Statistical Analysis": 9,
        "CHAPTER 4": 10,
        "4.1 Weight Loss": 10,
        "4.2 Surface Roughness": 11,
        "4.3 Summary of Results": 12,
        "CHAPTER 5": 13,
        "5.1 Weight Loss": 13,
        "5.2 Surface Roughness": 14,
        "5.3 Clinical Implications": 15,
        "5.4 Strengths of the Study": 15,
        "5.5 Limitations of the Study": 16,
        "5.6 Concluding Remarks": 16,
        "CHAPTER 6": 17,
        "6.1 Conclusions": 17,
        "6.2 Recommendations": 17,
        "CHAPTER 7": 19,
    }
    pretty_titles = {
        "INTRODUCTION": "Introduction",
        "REVIEW OF LITERATURE": "Review of Literature",
        "MATERIALS AND METHODS": "Materials and Methods",
        "RESULTS": "Results",
        "DISCUSSION": "Discussion",
        "CONCLUSIONS AND RECOMMENDATIONS": "Conclusions and Recommendations",
        "REFERENCES": "References",
    }
    toc_entries = []
    pending = None
    for raw in md.splitlines():
        line = raw.rstrip()
        if line.startswith("# CHAPTER"):
            pending = line[2:].strip()
        elif line.startswith("# ") and pending:
            title = line[2:].strip()
            pretty = pretty_titles.get(title, title)
            page = chapter_pages.get(pending, "")
            toc_entries.append(("ch", f"{pending}  {pretty}", page))
            pending = None
        elif line.startswith("## "):
            h = line[3:].strip()
            toc_entries.append(("sec", h, chapter_pages.get(h, "")))

    add_supervisors_page(doc)
    add_declaration(doc)
    add_abstract(doc)
    add_contents(doc, toc_entries)
    add_list_of_tables(doc)
    add_abbreviations(doc)

    # Body section — Arabic page numbers
    body_sect = doc.add_section()
    set_margins(body_sect)
    _setup_footer(body_sect, fmt="decimal", restart=1)

    convert_body(doc, md)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
